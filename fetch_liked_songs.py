# -*- coding: utf-8 -*-
"""抓取网易云音乐「我喜欢的音乐」歌单并导出 Excel/CSV"""
import urllib.request, urllib.parse, json, time, csv, sys

HEADERS = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64)', 'Referer': 'https://music.163.com/'}
PLAYLIST_ID = 497671214
OUT_PREFIX = r'C:\Users\Administrator\Desktop\WorkDir\mi-note-export\我喜欢的音乐_网易云'


def get(url, retries=3):
    for i in range(retries):
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            return json.load(urllib.request.urlopen(req, timeout=30))
        except Exception as e:
            print(f'  retry {i+1}: {e}', flush=True)
            time.sleep(2)
    raise RuntimeError('failed: ' + url)


# 1. 歌单信息 + 全部 trackIds（保持原始顺序）
d = get(f'https://music.163.com/api/v6/playlist/detail?id={PLAYLIST_ID}&n=1000')
pl = d['playlist']
tids = [t['id'] for t in pl['trackIds']]
print(f"歌单: {pl['name']}  共 {len(tids)} 首", flush=True)

# 2. 分批获取歌曲详情
BATCH = 200
songs = {}
for i in range(0, len(tids), BATCH):
    batch = tids[i:i+BATCH]
    url = 'https://music.163.com/api/song/detail?ids=' + urllib.parse.quote(json.dumps(batch))
    r = get(url)
    for s in r.get('songs', []):
        artists = '/'.join(a['name'] for a in s.get('artists', []))
        album = (s.get('album') or {}).get('name', '')
        dur = s.get('duration', 0)
        songs[s['id']] = {
            'song_id': s['id'],
            'name': s.get('name', ''),
            'artists': artists,
            'album': album,
            'duration': f"{dur//60000}:{dur//1000%60:02d}",
        }
    print(f'  已获取 {min(i+BATCH, len(tids))}/{len(tids)}', flush=True)
    time.sleep(0.3)

# 3. 按歌单原始顺序组装
rows = []
missing = 0
for idx, tid in enumerate(tids, 1):
    s = songs.get(tid)
    if s:
        rows.append([idx, s['name'], s['artists'], s['album'], s['duration'], s['song_id'],
                     f"https://music.163.com/song?id={s['song_id']}"])
    else:
        missing += 1
        rows.append([idx, '(获取失败)', '', '', '', tid, f"https://music.163.com/song?id={tid}"])

print(f'完成: {len(rows)} 行, 详情缺失 {missing} 首', flush=True)

# 4. 导出 CSV
csv_path = OUT_PREFIX + '.csv'
with open(csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.writer(f)
    w.writerow(['序号', '歌名', '歌手', '专辑', '时长', '歌曲ID', '链接'])
    w.writerows(rows)

# 5. 导出 Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook()
ws = wb.active
ws.title = '我喜欢的音乐'
header = ['序号', '歌名', '歌手', '专辑', '时长', '歌曲ID', '链接']
ws.append(header)
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='C20C0C')
    c.alignment = Alignment(horizontal='center')
for r in rows:
    ws.append(r)
for col, width in zip('ABCDEFG', [6, 40, 30, 30, 8, 12, 45]):
    ws.column_dimensions[col].width = width
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:G{len(rows)+1}'
xlsx_path = OUT_PREFIX + '.xlsx'
wb.save(xlsx_path)
print('CSV :', csv_path)
print('XLSX:', xlsx_path)
